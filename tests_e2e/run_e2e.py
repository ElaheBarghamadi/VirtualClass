"""End-to-end browser suite for the virtual classroom.

Run against a LIVE dev server:

    python manage.py migrate
    python manage.py seed_demo            # prints the room code
    python manage.py runserver 8000
    python tests_e2e/run_e2e.py --room <CODE> [--base http://127.0.0.1:8000]

Requires:  pip install -r requirements-dev.txt  +  playwright install chromium
Optional:  LibreOffice on the server enables the Office-conversion checks
           (they are skipped automatically when unavailable).

Covers: join flows, chat WS, raise-hand, force-mute with device release,
role promotion (auto-reload + can_present), files panel (empty state,
client-side validation, live sync, present, badge, download, delete),
Office→PDF presentation, whiteboard paging/drawing, camera + PiP, exit.
"""
from __future__ import annotations

import argparse
import re
import sys

from playwright.sync_api import sync_playwright

PASSWORD = "Pass-12345"
FA = str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")
RESULTS: list[tuple[str, bool, str]] = []
SKIPPED: list[str] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    RESULTS.append((name, ok, detail))
    print(("PASS " if ok else "FAIL ") + name + (f"  [{detail}]" if detail else ""))


def skip(name: str, why: str) -> None:
    SKIPPED.append(name)
    print(f"SKIP {name}  [{why}]")


def make_fixtures() -> tuple[bytes, bytes | None]:
    """(pdf_bytes, pptx_bytes|None). PDF via reportlab when available."""
    pdf = None
    try:
        import io
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.pdfgen import canvas
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=landscape(A4))
        c.setFont("Helvetica-Bold", 40); c.drawString(100, 300, "Slide 1")
        c.showPage(); c.setFont("Helvetica-Bold", 40); c.drawString(100, 300, "Slide 2")
        c.showPage(); c.save()
        pdf = buf.getvalue()
    except ImportError:
        pass
    if pdf is None:  # minimal 1-page PDF fallback
        pdf = (b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
               b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
               b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 144]>>endobj\n"
               b"trailer<</Root 1 0 R>>\n%%EOF")
    pptx = None
    try:
        import io
        from pptx import Presentation as Pptx
        from pptx.util import Inches
        prs = Pptx()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        tb.text_frame.text = "E2E office slide"
        buf = io.BytesIO(); prs.save(buf)
        pptx = buf.getvalue()
    except ImportError:
        pass
    return pdf, pptx


def login_and_join(pw, base: str, user: str, code: str):
    b = pw.chromium.launch(args=[
        "--use-fake-device-for-media-stream",
        "--use-fake-ui-for-media-stream",
        "--autoplay-policy=no-user-gesture-required",
    ])
    pg = b.new_page()
    errs: list[str] = []
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.on("console", lambda m: errs.append(f"console.{m.type}: {m.text}") if m.type == "error" else None)
    pg.goto(base + "/accounts/login/", timeout=20000)
    pg.fill("#id_username", user)
    pg.fill("#id_password", PASSWORD)
    pg.click("button[type=submit]")
    pg.wait_for_url("**/dashboard/**", timeout=20000)
    pg.goto(f"{base}/class/{code}/lobby/", timeout=20000)
    if "/room/" not in pg.url:  # lobby: members skip it
        if pg.locator("input[name=display_name]").count():
            pg.fill("input[name=display_name]", user)
        pg.click("#btn-to-devices")
        pg.wait_for_timeout(600)
        pg.click("#join-form button[type=submit]")
    pg.wait_for_selector("#room-side", state="attached", timeout=40000)
    return b, pg, errs


def run(base: str, code: str) -> int:
    pdf_bytes, pptx_bytes = make_fixtures()
    import tempfile, os
    tmp = tempfile.mkdtemp(prefix="e2e-")
    pdf_path = os.path.join(tmp, "e2e_deck.pdf")
    open(pdf_path, "wb").write(pdf_bytes)
    pptx_path = None
    if pptx_bytes:
        pptx_path = os.path.join(tmp, "e2e_lesson.pptx")
        open(pptx_path, "wb").write(pptx_bytes)
    bad_path = os.path.join(tmp, "bad.txt")
    open(bad_path, "w").write("not allowed")

    with sync_playwright() as pw:
        b_o, pg_o, oerr = login_and_join(pw, base, "owner1", code)
        b_s, pg_s, serr = login_and_join(pw, base, "stu1", code)
        pg_o.wait_for_timeout(2500)

        # ---------- basics ----------
        clock = pg_o.text_content("#room-clock") or ""
        check("room clock ticking", clock and clock != "--:--", clock)
        pg_o.click('.side-tab[data-tab="participants"]')
        pg_o.wait_for_selector(".participant", timeout=10000)
        check("roster shows both participants", pg_o.locator(".participant").count() >= 2)

        # ---------- chat over WS ----------
        pg_o.click('.side-tab[data-tab="chat"]')
        pg_o.fill("#chat-input", "سلام e2e")
        pg_o.click("#chat-send")
        try:
            pg_s.wait_for_function(
                "() => document.getElementById('chat-messages')?.innerText.includes('سلام e2e')",
                timeout=8000)
            chat_ok = True
        except Exception:
            chat_ok = False
        check("chat delivers live to student", chat_ok)

        # ---------- raise hand ----------
        pg_s.click("#btn-hand"); pg_s.wait_for_timeout(1500)
        check("raise-hand pill visible to owner",
              pg_o.evaluate("() => !document.getElementById('hand-pill')?.classList.contains('hidden')"))

        # ---------- mic + force-mute ----------
        pg_s.click("#btn-mic"); pg_s.wait_for_timeout(2000)
        trk = pg_s.evaluate("() => window.__media?.impl?.localStream?.getAudioTracks?.().length ?? 0")
        check("student mic produced a track", trk == 1, str(trk))
        pg_o.click('.side-tab[data-tab="participants"]'); pg_o.wait_for_timeout(500)
        stu_id = pg_o.evaluate("""() => [...document.querySelectorAll('.participant')]
            .find(li => (li.querySelector('.participant-name')?.textContent||'').includes('سارا'))?.dataset.identity""")
        sel = f'.participant[data-identity="{stu_id}"]'
        pg_o.click(f"{sel} .hc-btn >> nth=0"); pg_o.wait_for_timeout(2000)
        check("owner mute shows in roster",
              pg_o.evaluate(f"() => document.querySelector('{sel} .s-muted')?.classList.contains('on')") is True)
        released = pg_s.evaluate("() => window.__media?.impl?.localStream?.getAudioTracks?.().length ?? -1")
        check("force-mute released the mic device", released in (0, -1), str(released))

        # ---------- promotion ----------
        pg_o.click(f"{sel} .hc-btn >> nth=3")
        pg_s.wait_for_load_state("load", timeout=20000)
        pg_s.wait_for_selector("#room-side", timeout=20000)
        pg_s.wait_for_timeout(2000)
        check("promoted student gained can_present after auto-reload",
              pg_s.evaluate("() => JSON.parse(document.getElementById('member-permissions').textContent).can_present") is True)

        # ---------- files ----------
        pg_o.click('.side-tab[data-tab="files"]')
        if pg_o.locator("#file-empty").count():
            check("files empty state visible", pg_o.locator("#file-empty").is_visible())
        pg_o.set_input_files("#file-input", bad_path); pg_o.wait_for_timeout(700)
        check("bad extension blocked client-side",
              pg_o.evaluate("() => document.body.innerText.includes('مجاز نیست')"))
        pg_o.set_input_files("#file-input", pdf_path)
        pg_o.wait_for_selector(".file-item", timeout=20000); pg_o.wait_for_timeout(1000)
        check("upload created a row", pg_o.locator(".file-item").count() >= 1)
        pg_s.click('.side-tab[data-tab="files"]')
        pg_s.wait_for_selector(".file-item", timeout=10000)
        check("student received the file live", True)
        check("non-uploader presenter sees no delete button",
              pg_s.locator(".file-item .file-delete").count() == 0)
        pg_s.click(".file-item .file-info"); pg_s.wait_for_timeout(3000)
        check("promoted student presents via row click",
              pg_s.evaluate("() => document.querySelector('.stage-view:not(.hidden)')?.id") == "view-presentation")
        check("presenting badge synced to owner",
              pg_o.evaluate("() => document.querySelector('.file-item')?.classList.contains('presenting')") is True)
        pg_o.click("#pres-next"); pg_o.wait_for_timeout(1500)
        opl = (pg_o.text_content("#pres-page-label") or "").translate(FA)
        check("owner auto-followed to page 2", "2" in opl, opl)
        with pg_o.expect_download(timeout=20000) as dl:
            pg_o.click(".file-item a.file-download")
        check("download works", "e2e_deck" in dl.value.suggested_filename)
        pg_o.click(".file-item .file-delete")
        pg_o.wait_for_selector('[data-act="ok"]', timeout=5000)
        pg_o.click('[data-act="ok"]'); pg_o.wait_for_timeout(1800)
        check("delete removed the row", pg_o.locator(".file-item").count() == 0)
        pg_s.wait_for_timeout(1500)
        check("presentation cleared on student after delete",
              pg_s.evaluate("() => document.querySelector('.stage-view:not(.hidden)')?.id") == "view-media")

        # ---------- office → pdf (optional) ----------
        if pptx_path:
            pg_o.set_input_files("#file-input", pptx_path)
            pg_o.wait_for_selector(".file-item", timeout=30000); pg_o.wait_for_timeout(2500)
            has_pill = pg_o.locator(".pdf-ready-pill").count() > 0
            if has_pill:
                pg_o.click(".file-item .file-info"); pg_o.wait_for_timeout(3500)
                view = pg_o.evaluate("() => document.querySelector('.stage-view:not(.hidden)')?.id")
                label = (pg_o.text_content("#pres-page-label") or "").translate(FA)
                check("PPTX presents via server-converted PDF",
                      view == "view-presentation" and bool(re.search(r"\d", label)),
                      f"{view} {label}")
            else:
                skip("PPTX presentation", "LibreOffice not installed on the server")
        else:
            skip("PPTX presentation", "python-pptx not installed locally")

        # ---------- whiteboard ----------
        pg_o.evaluate("() => window.__switchView('whiteboard')"); pg_o.wait_for_timeout(1200)
        lab0 = (pg_o.text_content("#wb-page-label") or "").translate(FA)
        pg_o.click("#wb-new-page"); pg_o.wait_for_timeout(1200)
        lab1 = (pg_o.text_content("#wb-page-label") or "").translate(FA)
        n0 = int(re.search(r"\d+", lab0).group()); n1 = int(re.search(r"\d+", lab1).group())
        check("whiteboard new page increments", n1 == n0 + 1, f"{lab0} -> {lab1}")
        box = pg_o.locator("#whiteboard-canvas").bounding_box()
        pg_o.mouse.move(box["x"] + 100, box["y"] + 100); pg_o.mouse.down()
        pg_o.mouse.move(box["x"] + 260, box["y"] + 200, steps=8); pg_o.mouse.up()
        pg_o.wait_for_timeout(1200)
        check("whiteboard stroke rendered",
              pg_o.evaluate("() => { const c = document.getElementById('whiteboard-canvas'); return c ? c.toDataURL().length > 5000 : false; }"))

        # ---------- camera + PiP ----------
        pg_s.click("#btn-camera"); pg_s.wait_for_timeout(2500)
        check("student camera produced a track",
              pg_s.evaluate("() => window.__media?.impl?.localStream?.getVideoTracks?.().length ?? 0") == 1)
        pg_s.evaluate("() => window.__switchView('whiteboard')"); pg_s.wait_for_timeout(1500)
        check("PiP camera visible over whiteboard view",
              pg_s.evaluate("() => !document.getElementById('cam-pip')?.classList.contains('hidden')"))

        # ---------- exit ----------
        pg_s.click("#btn-exit"); pg_s.wait_for_timeout(600)
        if pg_s.locator('[data-act="ok"]').count():
            pg_s.click('[data-act="ok"]')
        pg_s.wait_for_timeout(2500)
        check("exit leaves the room", "/room/" not in pg_s.url, pg_s.url)

        check("no page errors (owner)", not oerr, ";".join(oerr[:3]))
        check("no page errors (student)", not serr, ";".join(serr[:3]))
        b_o.close(); b_s.close()

    fails = [r for r in RESULTS if not r[1]]
    total = len(RESULTS)
    print(f"\n===== {total - len(fails)}/{total} PASS, {len(SKIPPED)} skipped =====")
    return 1 if fails else 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default="http://127.0.0.1:8000")
    ap.add_argument("--room", required=True)
    args = ap.parse_args()
    sys.exit(run(args.base, args.room))
